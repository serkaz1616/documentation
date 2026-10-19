#!/usr/bin/env python3
"""
Lidl-Filialdaten (Öffnungszeiten, Adresse, Ausstattung) direkt von lidl.de holen
und in die Filial-Excel schreiben.

Zwei Schritte, damit später neue Spalten ohne erneutes Abrufen ergänzt werden können:

  1) scrape  – liest alle Filialseiten aus der offiziellen Lidl-Sitemap und speichert
               die Rohdaten je Filiale in einer JSON-Datei (Cache, fortsetzbar).
  2) merge   – liest die Excel, ordnet über die Filialnummer zu (668 -> DE00668)
               und schreibt eine neue Excel mit zusätzlichen Spalten.

Beispiele:
  python lidl_filialen.py scrape
  python lidl_filialen.py merge --input Filialen_DE.xlsx --output Filialen_DE_lidl.xlsx
  python lidl_filialen.py all   --input Filialen_DE.xlsx --output Filialen_DE_lidl.xlsx

Quelle ist ausschließlich lidl.de (keine Google-/Fremddaten).
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import threading
import time
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SITEMAP_URL = "https://www.lidl.de/s/de-DE/filialen/sitemap.xml"
STORE_URL_RE = re.compile(r"^https://www\.lidl\.de/s/de-DE/filialen/[^/]+/[^/]+/$")
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"
)
DEFAULT_CACHE = Path(__file__).with_name("lidl_filialen_cache.json")

DAYS = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
DAY_DE = {"monday": "Mo", "tuesday": "Di", "wednesday": "Mi", "thursday": "Do",
          "friday": "Fr", "saturday": "Sa", "sunday": "So"}


# --------------------------------------------------------------------------- #
# Abruf
# --------------------------------------------------------------------------- #

_thread_local = threading.local()


def session() -> requests.Session:
    s = getattr(_thread_local, "session", None)
    if s is None:
        s = requests.Session()
        s.headers.update({"User-Agent": USER_AGENT, "Accept-Language": "de-DE,de;q=0.9"})
        _thread_local.session = s
    return s


def get(url: str, retries: int = 4, timeout: int = 30) -> requests.Response:
    last_exc: Exception | None = None
    for attempt in range(retries):
        try:
            r = session().get(url, timeout=timeout)
            if r.status_code == 200:
                return r
            if r.status_code == 404:
                r.raise_for_status()
            last_exc = requests.HTTPError(f"HTTP {r.status_code} für {url}")
        except requests.HTTPError:
            raise
        except requests.RequestException as e:
            last_exc = e
        time.sleep(2 ** (attempt + 1))
    raise last_exc  # type: ignore[misc]


def store_urls_from_sitemap() -> list[str]:
    xml = get(SITEMAP_URL).content
    ns = {"sm": "http://www.sitemaps.org/schemas/sitemap/0.9"}
    locs = [e.text.strip() for e in ET.fromstring(xml).findall(".//sm:loc", ns) if e.text]
    return sorted({u for u in locs if STORE_URL_RE.match(u)})


def decode_nuxt_payload(data: list):
    """Nuxt-Payload (devalue-Format) in normale Python-Objekte umwandeln."""
    memo: dict[int, object] = {}
    specials = {-1: None, -2: None, -3: float("nan"), -4: float("inf"), -5: float("-inf"), -6: 0}

    def resolve(i):
        if isinstance(i, int) and i < 0:
            return specials.get(i)
        if i in memo:
            return memo[i]
        v = data[i]
        if isinstance(v, dict):
            out: dict = {}
            memo[i] = out
            for k, idx in v.items():
                out[k] = resolve(idx)
            return out
        if isinstance(v, list):
            if v and isinstance(v[0], str):
                tag = v[0]
                if tag in ("Reactive", "ShallowReactive", "Ref", "ShallowRef"):
                    res = resolve(v[1])
                    memo[i] = res
                    return res
                if tag == "EmptyRef" or tag == "EmptyShallowRef":
                    return None
                if tag == "Date":
                    return v[1]
                if tag in ("Set",):
                    return [resolve(x) for x in v[1:]]
                if tag == "Map":
                    return {str(resolve(v[j])): resolve(v[j + 1]) for j in range(1, len(v), 2)}
                if tag == "null":
                    return None
            out_l: list = []
            memo[i] = out_l
            out_l.extend(resolve(x) for x in v)
            return out_l
        return v

    return resolve(0)


def find_store(obj):
    """Sucht im dekodierten Payload das Filial-Objekt (mit objectNumber + Öffnungszeiten)."""
    stack = [obj]
    seen = set()
    while stack:
        cur = stack.pop()
        if id(cur) in seen:
            continue
        seen.add(id(cur))
        if isinstance(cur, dict):
            if "objectNumber" in cur and "generalOpeningHours" in cur:
                return cur
            stack.extend(cur.values())
        elif isinstance(cur, list):
            stack.extend(cur)
    return None


def fetch_store(url: str) -> dict:
    r = get(url + "_payload.json")
    store = find_store(decode_nuxt_payload(r.json()))
    if store is None:
        raise ValueError(f"Keine Filialdaten gefunden: {url}")
    store = dict(store)
    store.pop("nearbyStores", None)   # nicht benötigt, macht den Cache nur groß
    store.pop("openingHours", None)   # nur die nächsten 7 Tage; Regelzeiten stehen in generalOpeningHours
    store["_url"] = url
    store["_abgerufen"] = dt.datetime.now().isoformat(timespec="seconds")
    return store


def load_cache(path: Path) -> dict:
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {"stores": {}, "errors": {}}


def save_cache(path: Path, cache: dict) -> None:
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(path)


def cmd_scrape(args) -> None:
    cache_path = Path(args.cache)
    cache = load_cache(cache_path)
    urls = store_urls_from_sitemap()
    print(f"Sitemap: {len(urls)} Filialseiten")

    if args.refresh:
        done_urls: set[str] = set()
        cache = {"stores": {}, "errors": {}}
    else:
        done_urls = {s["_url"] for s in cache["stores"].values()}
    todo = [u for u in urls if u not in done_urls]
    if args.limit:
        todo = todo[: args.limit]
    print(f"Abzurufen: {len(todo)} (bereits im Cache: {len(done_urls)})")

    lock = threading.Lock()
    cache["errors"] = {}

    def work(u):
        time.sleep(args.delay)
        return u, fetch_store(u)

    n = 0
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futures = [ex.submit(work, u) for u in todo]
        for f in as_completed(futures):
            n += 1
            try:
                u, st = f.result()
                with lock:
                    cache["stores"][st["objectNumber"]] = st
            except Exception as e:  # noqa: BLE001
                with lock:
                    cache["errors"][str(e)[:200]] = time.time()
            if n % 100 == 0 or n == len(todo):
                with lock:
                    save_cache(cache_path, cache)
                print(f"  {n}/{len(todo)}  Fehler: {len(cache['errors'])}", flush=True)

    cache["sitemap_urls"] = len(urls)
    cache["stand"] = dt.datetime.now().isoformat(timespec="seconds")
    save_cache(cache_path, cache)
    print(f"Fertig. {len(cache['stores'])} Filialen im Cache: {cache_path}")
    if cache["errors"]:
        print(f"{len(cache['errors'])} Fehler – einfach erneut 'scrape' starten, es wird fortgesetzt.")


# --------------------------------------------------------------------------- #
# Aufbereitung
# --------------------------------------------------------------------------- #

def fmt_ranges(ranges: list[dict]) -> str:
    if not ranges:
        return "geschlossen"
    return ", ".join(f"{r['from']}-{r['to']}" for r in ranges)


def opening_hours_summary(regular: dict) -> str:
    """z. B. 'Mo-Sa 07:00-22:00; So geschlossen'."""
    parts = []
    i = 0
    while i < 7:
        val = fmt_ranges(regular.get(DAYS[i]) or [])
        j = i
        while j + 1 < 7 and fmt_ranges(regular.get(DAYS[j + 1]) or []) == val:
            j += 1
        days = DAY_DE[DAYS[i]] if i == j else f"{DAY_DE[DAYS[i]]}-{DAY_DE[DAYS[j]]}"
        parts.append(f"{days} {val}")
        i = j + 1
    return "; ".join(parts)


def special_summary(special: list[dict]) -> str:
    out = []
    for s in special or []:
        d = s.get("date", "")
        try:
            d = dt.date.fromisoformat(d).strftime("%d.%m.%Y")
        except ValueError:
            pass
        out.append(f"{d} {fmt_ranges(s.get('timeRanges') or [])}")
    return "; ".join(out)


def norm_street(s: str | None) -> str:
    s = (s or "").lower().strip()
    s = s.replace("ß", "ss").replace("ä", "ae").replace("ö", "oe").replace("ü", "ue")
    s = re.sub(r"strasse\b|str\b\.?", "str", s)
    s = re.sub(r"[^a-z0-9]", "", s)
    return s


def norm_city(s: str | None) -> str:
    s = (s or "").lower()
    s = s.replace("ß", "ss").replace("ä", "ae").replace("ö", "oe").replace("ü", "ue")
    return re.sub(r"[^a-z]", "", s)


def address_check(plz, ort, strasse, st: dict) -> str:
    a = st.get("address") or {}
    issues = []
    if str(plz or "").strip().zfill(5) != str(a.get("zip") or "").strip().zfill(5):
        issues.append("PLZ")
    lidl_street = " ".join(x for x in [a.get("streetName"), a.get("streetNumber")] if x)
    if norm_street(strasse) != norm_street(lidl_street):
        issues.append("Straße")
    # Excel-Ort enthält oft den Stadtteil ("Mannheim-Tattersall"), Lidl nur die Stadt
    city_x, city_l = norm_city(ort), norm_city(a.get("city"))
    store_name = norm_city(st.get("storeName"))
    if not (city_l and (city_l in city_x or city_x in city_l or city_x == store_name)):
        issues.append("Ort")
    return "OK" if not issues else "Abweichung: " + ", ".join(issues)


def lidl_columns(st: dict | None, plz=None, ort=None, strasse=None) -> dict:
    if st is None:
        return {"Lidl gefunden": "nein"}
    a = st.get("address") or {}
    goh = st.get("generalOpeningHours") or {}
    regular = goh.get("regular") or {}
    md = st.get("marketingData") or {}
    icons = md.get("infoIcons") or []
    cols = {
        "Lidl gefunden": "ja",
        "Lidl Objektnummer": st.get("objectNumber"),
        "Lidl Filialname": st.get("storeName"),
        "Lidl Straße": " ".join(x for x in [a.get("streetName"), a.get("streetNumber")] if x),
        "Lidl PLZ": a.get("zip"),
        "Lidl Ort": a.get("city"),
        "Lidl Bundesland": a.get("state"),
        "Breitengrad": a.get("latitude"),
        "Längengrad": a.get("longitude"),
        "Adressabgleich": address_check(plz, ort, strasse, st) if plz is not None else "",
        "Öffnungszeiten (Lidl)": opening_hours_summary(regular),
    }
    for d in DAYS:
        cols[f"Öffnung {DAY_DE[d]}"] = fmt_ranges(regular.get(d) or [])
    cols.update({
        "Sonderöffnungszeiten": special_summary(goh.get("special") or []),
        "Ausstattung": ", ".join(i.get("name", "") for i in icons),
        "Vertriebsgesellschaft (Lidl)": md.get("offerRegionName"),
        "Lidl URL": st.get("_url"),
        "Abgerufen am": st.get("_abgerufen"),
    })
    return cols


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #

HEADER_FILL = PatternFill("solid", fgColor="FFF2CC")
WARN_FILL = PatternFill("solid", fgColor="F8CBAD")


def find_col(headers: list, *names) -> int | None:
    norm = [str(h or "").strip().lower() for h in headers]
    for n in names:
        if n.lower() in norm:
            return norm.index(n.lower())
    return None


def to_object_number(filialnr) -> str | None:
    if filialnr is None or str(filialnr).strip() == "":
        return None
    digits = re.sub(r"\D", "", str(filialnr))
    return f"DE{int(digits):05d}" if digits else None


def autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), 60)


def cmd_merge(args) -> None:
    cache = load_cache(Path(args.cache))
    stores: dict = cache["stores"]
    if not stores:
        sys.exit("Cache ist leer – zuerst 'scrape' ausführen.")

    wb = load_workbook(args.input)
    ws = wb[args.sheet] if args.sheet else wb.active
    headers = [c.value for c in ws[1]]
    c_nr = find_col(headers, "Filialnummer", "Filiale", "Filial-Nr", "Filialnr")
    c_plz = find_col(headers, "PLZ")
    c_ort = find_col(headers, "Ort")
    c_str = find_col(headers, "Straße", "Strasse")
    if c_nr is None:
        sys.exit(f"Spalte 'Filialnummer' nicht gefunden. Vorhandene Spalten: {headers}")

    # Bereits vorhandene Lidl-Spalten bei erneutem Lauf überschreiben statt doppelt anhängen
    template = list(lidl_columns({"_url": ""}, 0, "", "").keys())
    first_new = len(headers) + 1
    for name in template:
        if name in headers:
            first_new = min(first_new, headers.index(name) + 1)
    if first_new <= len(headers):
        ws.delete_cols(first_new, len(headers) - first_new + 1)

    for k, name in enumerate(template):
        cell = ws.cell(row=1, column=first_new + k, value=name)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL

    used = set()
    stats = {"gefunden": 0, "nicht gefunden": 0, "Adresse OK": 0, "Adresse abweichend": 0}
    for row in ws.iter_rows(min_row=2):
        nr = row[c_nr].value
        if nr is None:
            continue
        on = to_object_number(nr)
        st = stores.get(on) if on else None
        vals = lidl_columns(
            st,
            row[c_plz].value if c_plz is not None else None,
            row[c_ort].value if c_ort is not None else None,
            row[c_str].value if c_str is not None else None,
        )
        if st:
            used.add(on)
            stats["gefunden"] += 1
            if vals["Adressabgleich"] == "OK":
                stats["Adresse OK"] += 1
            elif vals["Adressabgleich"]:
                stats["Adresse abweichend"] += 1
        else:
            stats["nicht gefunden"] += 1
        r = row[0].row
        for k, name in enumerate(template):
            cell = ws.cell(row=r, column=first_new + k, value=vals.get(name))
            if (name == "Adressabgleich" and str(cell.value).startswith("Abweichung")) or \
               (name == "Lidl gefunden" and cell.value == "nein"):
                cell.fill = WARN_FILL

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    # Zusatzblatt: Filialen auf lidl.de, die nicht in der Excel stehen
    name = "Nur auf lidl.de"
    if name in wb.sheetnames:
        del wb[name]
    ws2 = wb.create_sheet(name)
    hdr = ["Filialnummer"] + template
    ws2.append(hdr)
    for c in ws2[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
    for on in sorted(set(stores) - used):
        vals = lidl_columns(stores[on])
        ws2.append([int(on[2:])] + [vals.get(h) for h in template])
    ws2.freeze_panes = "B2"
    ws2.auto_filter.ref = ws2.dimensions
    autosize(ws2)

    # Info-Blatt
    name = "Info Lidl-Abgleich"
    if name in wb.sheetnames:
        del wb[name]
    ws3 = wb.create_sheet(name)
    ws3.append(["Quelle", "lidl.de Filialseiten (Sitemap " + SITEMAP_URL + ")"])
    ws3.append(["Stand Abruf", cache.get("stand")])
    ws3.append(["Filialen auf lidl.de", len(stores)])
    for k, v in stats.items():
        ws3.append([f"Excel-Zeilen: {k}", v])
    ws3.append(["Nur auf lidl.de (nicht in Excel)", len(set(stores) - used)])
    autosize(ws3)

    wb.save(args.output)
    print(f"Gespeichert: {args.output}")
    for k, v in stats.items():
        print(f"  {k}: {v}")
    print(f"  nur auf lidl.de: {len(set(stores) - used)}")


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--cache", default=str(DEFAULT_CACHE), help="JSON-Cache mit Rohdaten")
    sub = p.add_subparsers(dest="cmd", required=True)

    def scrape_opts(sp):
        sp.add_argument("--workers", type=int, default=4, help="parallele Abrufe (Standard 4)")
        sp.add_argument("--delay", type=float, default=0.5, help="Pause je Abruf in s (Standard 0.5)")
        sp.add_argument("--limit", type=int, default=0, help="nur N Filialen (zum Testen)")
        sp.add_argument("--refresh", action="store_true", help="Cache verwerfen und alles neu laden")

    def merge_opts(sp):
        sp.add_argument("--input", required=True, help="Excel mit Filialnummern")
        sp.add_argument("--output", required=True, help="Ziel-Excel")
        sp.add_argument("--sheet", default=None, help="Tabellenblatt (Standard: erstes)")

    scrape_opts(sub.add_parser("scrape", help="Filialdaten von lidl.de laden"))
    merge_opts(sub.add_parser("merge", help="Daten in Excel schreiben"))
    sp_all = sub.add_parser("all", help="scrape + merge")
    scrape_opts(sp_all)
    merge_opts(sp_all)

    args = p.parse_args()
    if args.cmd in ("scrape", "all"):
        cmd_scrape(args)
    if args.cmd in ("merge", "all"):
        cmd_merge(args)


if __name__ == "__main__":
    main()
